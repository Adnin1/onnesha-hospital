import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern Medical Indigo & Teal
    C_PRIMARY = RGBColor(13, 71, 161)      # Deep Navy/Indigo
    C_SECONDARY = RGBColor(0, 150, 136)    # Medical Teal
    C_ACCENT = RGBColor(255, 111, 0)       # Amber/Orange
    C_DARK = RGBColor(33, 33, 33)          # Dark Charcoal
    C_WHITE = RGBColor(255, 255, 255)      # Pure White
    C_LIGHT_BG = RGBColor(245, 247, 250)   # Light Gray/Slate
    C_CARD_BG = RGBColor(255, 255, 255)    # Card White
    C_CARD_BORDER = RGBColor(207, 216, 220)
    C_EMERALD = RGBColor(46, 125, 50)      # Success Green
    C_CRIMSON = RGBColor(198, 40, 40)      # Emergency Red

    def add_header(slide, title_text, subtitle_text, step_badge=None):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = C_PRIMARY
        banner.line.color.rgb = C_PRIMARY

        # Accent bottom line
        accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.08))
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = C_SECONDARY
        accent_line.line.color.rgb = C_SECONDARY

        # Title text
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(9.5), Inches(0.55))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.font.name = "Segoe UI"

        # Subtitle text
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(13)
        p2.font.color.rgb = RGBColor(200, 230, 240)
        p2.font.name = "Segoe UI"

        # Badge if provided
        if step_badge:
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.25), Inches(2.2), Inches(0.65))
            badge.fill.solid()
            badge.fill.fore_color.rgb = C_ACCENT
            badge.line.color.rgb = C_ACCENT
            btf = badge.text_frame
            bp = btf.paragraphs[0]
            bp.text = step_badge
            bp.alignment = PP_ALIGN.CENTER
            bp.font.size = Pt(14)
            bp.font.bold = True
            bp.font.color.rgb = C_WHITE
            bp.font.name = "Segoe UI"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_LIGHT_BG
        bg.line.color.rgb = C_LIGHT_BG
        return bg

    def add_card(slide, left, top, width, height, title, items, badge_text=None, border_color=C_CARD_BORDER, header_bg=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        # Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = C_PRIMARY
        p_title.font.name = "Segoe UI"

        if badge_text:
            p_badge = tf.add_paragraph()
            p_badge.text = f"[{badge_text}]"
            p_badge.font.size = Pt(11)
            p_badge.font.bold = True
            p_badge.font.color.rgb = C_SECONDARY
            p_badge.font.name = "Segoe UI"

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"- {item}"
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = C_DARK
            p_item.font.name = "Segoe UI"
            p_item.space_before = Pt(4)

        return card

    # ==========================================
    # SLIDE 1: Title Slide (Cover)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_PRIMARY
    bg1.line.color.rgb = C_PRIMARY

    # Title box
    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.8))
    ttf = tbox.text_frame
    ttf.word_wrap = True

    p = ttf.paragraphs[0]
    p.text = "ONNESHA HOSPITAL MANAGEMENT SYSTEM"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER

    p2 = ttf.add_paragraph()
    p2.text = "রিসেপশন এ-টু-জেড সহজ নির্দেশিকা (Master Reception Guide)"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(128, 222, 234)
    p2.font.name = "Segoe UI"
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(14)

    p3 = ttf.add_paragraph()
    p3.text = "ছোট বাচ্চাদের মতো সহজ ধাপে সম্পূর্ণ হাসপাতালের রিসেপশন ও কাউন্টার পরিচালনা"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(224, 242, 241)
    p3.font.name = "Segoe UI"
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(10)

    p4 = ttf.add_paragraph()
    p4.text = "লাইভ সিস্টেম: https://onnesha-hospital.pages.dev  |  সংস্করণ: ১.১.৯০  |  সম্পূর্ণ নির্ভুল ও নিরাপদ"
    p4.font.size = Pt(13)
    p4.font.color.rgb = RGBColor(255, 213, 79)
    p4.font.name = "Segoe UI"
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(24)

    # ==========================================
    # SLIDE 2: Step 1 - Morning Login & Shift Start
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "১. সকালের লগইন ও শিফট শুরু (Login & Opening Cash)", "কম্পিউটার চালু করে রিসেপশন শুরু করার প্রথম ৩টি সহজ ধাপ", "ধাপ ০১")

    add_card(slide2, 0.8, 1.5, 3.6, 5.3, "ধাপ ১: ব্রাউজার ও লিঙ্ক ওপেন", [
        "কম্পিউটারে Google Chrome বা Edge খুলুন।",
        "লিঙ্ক টাইপ করুন: onnesha-hospital.pages.dev",
        "লগইন স্ক্রিন আসলে নিজের ইউজারনেম ও পাসওয়ার্ড দিন।",
        "কাউন্টার সিলেক্ট করুন (কাউন্টার ১ বা ২)।",
        "লগইন বাটনে ক্লিক করলেই মূল ড্যাশবোর্ড চলে আসবে।"
    ], "Login Protocol")

    add_card(slide2, 4.8, 1.5, 3.6, 5.3, "ধাপ ২: ওপেনিং ক্যাশ ড্রয়ার ইনপুট", [
        "দিনের শুরুতে ক্যাশ ড্রয়ারে খুচরা টাকা থাকে (যেমন ৳ ৫০০ বা ৳ ১,০০০)।",
        "সিস্টেমে 'Open Shift' বাটন দেখতে পাবেন।",
        "ড্রয়ারে যত টাকা আছে সেই এমাউন্টটি লিখে 'Start Shift' চাপুন।",
        "এই টাকা আপনার মূল হিসাবের সাথে অটোমেটিক যুক্ত থাকবে।",
        "সারাদিনের হিসাব স্বয়ংক্রিয়ভাবে মিলবে।"
    ], "Cash Opening")

    add_card(slide2, 8.8, 1.5, 3.7, 5.3, "গোল্ডেন বেবি রুলস (সহজ নিয়ম)", [
        "নিয়ম ১: নিজের পাসওয়ার্ড কাউকে বলবেন না।",
        "নিয়ম ২: কাউন্টার ছাড়ার সময় স্ক্রিন লক করে যান।",
        "নিয়ম ৩: প্রিন্টারের পাওয়ার ও পেপার রোল চেক করে নিন।",
        "প্রিন্টারে পর্যাপ্ত ৮০ মিমি পেপার আছে কিনা নিশ্চিত করুন।",
        "ইন্টারনেট কানেকশন ঠিক থাকলে উপরে সবুজ আইকন দেখাবে।"
    ], "Baby Safety", border_color=C_EMERALD)

    # ==========================================
    # SLIDE 3: Step 2 - Patient Registration
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "২. নতুন রোগী নিবন্ধন (New Patient Registration)", "হাসপাতালে নতুন রোগী আসলে মাত্র ১ মিনিটে নিবন্ধন করার নিয়ম", "ধাপ ০২")

    add_card(slide3, 0.8, 1.5, 3.6, 5.3, "ধাপ ১: 'New Patient' বাটন", [
        "স্ক্রিনের উপরের বামে '+ New Patient' বাটনে ক্লিক করুন।",
        "কীবোর্ড শর্টকাট: চাপুন Alt + N",
        "সাথে সাথে একটি পরিষ্কার সহজ ফর্ম স্ক্রিনে ভেসে উঠবে।",
        "রোগী পুরাতন হলে শুধু মোবাইল নম্বর দিয়ে সার্চ করুন।"
    ], "Form Open")

    add_card(slide3, 4.8, 1.5, 3.6, 5.3, "ধাপ ২: রোগীর তথ্য পূরণ", [
        "১. রোগীর পুরো নাম লিখুন (বাংলা বা ইংরেজি)।",
        "২. মোবাইল নম্বর (১১ ডিজিট বাধ্যতামূলক - এসএমএস যাবে)।",
        "৩. বয়স (যেমন ২৫ বছর) ও লিঙ্গ (পুরুষ / মহিলা / অন্যান্য)।",
        "৪. রক্তের গ্রুপ (জানা থাকলে নির্বাচন করুন)।",
        "৫. ঠিকানা / জেলা নির্বাচন করুন।"
    ], "Information Input")

    add_card(slide3, 8.8, 1.5, 3.7, 5.3, "ধাপ ৩: সংরক্ষণ ও আইডি তৈরি", [
        "'Save Patient' বাটনে ক্লিক করুন।",
        "সিস্টেম সাথে সাথে একটি ইউনিক Patient ID বানাবে (যেমনঃ OH-P-10025)।",
        "রোগীর মোবাইলে সাথে সাথে কনফার্মেশন এসএমএস চলে যাবে।",
        "এখন রোগী যে কোনো ডাক্তার বা টেস্টের জন্য রেডি।"
    ], "Unique ID Generation", border_color=C_SECONDARY)

    # ==========================================
    # SLIDE 4: Step 3 - Doctor OPD Appointment
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "৩. ডাক্তার অ্যাপয়েন্টমেন্ট ও টিকিট বুকিং (OPD Booking)", "রোগী কোন ডাক্তারের কাছে যাবে তা খুব সহজে সিলেক্ট ও বুকিং করার নিয়ম", "ধাপ ০৩")

    add_card(slide4, 0.8, 1.5, 3.6, 5.3, "ধাপ ১: ডিপার্টমেন্ট ও ডাক্তার নির্বাচন", [
        "রোগীকে জিজ্ঞেস করুন তার সমস্যা কি।",
        "ডিপার্টমেন্ট লিস্ট থেকে বিভাগ বেছে নিন (যেমন মেডিসিন, হৃদরোগ, শিশু, গাইনি)।",
        "আজকে উপস্থিত ডাক্তারদের তালিকা স্ক্রিনে দেখতে পাবেন।",
        "যে ডাক্তার রোগী দেখাতে চান তার নামের পাশে ক্লিক করুন।"
    ], "Doctor Select")

    add_card(slide4, 4.8, 1.5, 3.6, 5.3, "ধাপ ২: শিফট ও সময় নির্বাচন", [
        "ডাক্তার সাহেবের শিফট সিলেক্ট করুন (সকাল বা বিকাল)।",
        "সিস্টেম আপনাকে স্বয়ংক্রিয় সিরিয়াল নম্বর দেখাবে (যেমন সিরিয়াল # ০৭)।",
        "ভিজিট ফি অটোমেটিক স্ক্রিনে চলে আসবে (যেমন ৳ ৬০০ বা ৳ ৮০০)।",
        "কোনো রেফারাল ডাক্তার থাকলে ড্রপডাউন থেকে সিলেক্ট করতে পারেন।"
    ], "Slot & Fee")

    add_card(slide4, 8.8, 1.5, 3.7, 5.3, "ধাপ ৩: টিকিট নিশ্চিতকরণ", [
        "'Confirm Appointment' বাটনে ক্লিক করুন।",
        "সাথে সাথে থার্মাল প্রিন্টারে প্রেসক্রিপশন ও টোকেন স্লিপ প্রিন্ট হবে।",
        "রোগীকে হাসিমুখে বলুন: 'আপনার সিরিয়াল ৭, রুম নাম্বার ২০২ এ গিয়ে বসুন'।",
        "ডাক্তার সাহেবের কম্পিউটারেও রোগীর নাম অটোমেটিক চলে যাবে।"
    ], "Token Finalize", border_color=C_EMERALD)

    # ==========================================
    # SLIDE 5: Step 4 - Thermal 80mm Token Print
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "৪. ৮০ মিমি থার্মাল টোকেন প্রিন্ট (Thermal Receipt Print)", "রোগীর হাতে টোকেন স্লিপ প্রিন্ট করে দেওয়ার সহজ পদ্ধতি", "ধাপ ০৪")

    add_card(slide5, 0.8, 1.5, 3.6, 5.3, "স্লিপে যা যা স্বয়ংক্রিয় ছাপা হবে", [
        "হাসপাতালের নাম ও লোগো (Onnesha Hospital)।",
        "টোকেন ও সিরিয়াল নম্বর (বড় অক্ষরে স্পষ্ট দেখা যাবে)।",
        "রোগীর নাম, বয়স, জেন্ডার ও মোবাইল নম্বর।",
        "ডাক্তারের নাম, স্পেশালিটি ও রুম নম্বর (যেমন Room: 204)।",
        "তারিখ ও সময় এবং বারকোড (যা দিয়ে ডাক্তার স্ক্যান করতে পারেন)।"
    ], "Receipt Content")

    add_card(slide5, 4.8, 1.5, 3.6, 5.3, "১-ক্লিক প্রিন্ট করার উপায়", [
        "বুকিং শেষ হওয়ামাত্র প্রিন্ট ডায়ালগ স্বয়ংক্রিয় ওপেন হবে।",
        "কীবোর্ড শর্টকাট: Ctrl + P",
        "এন্টার (Enter) চাপলেই সাথে সাথে ৮০ মিমি প্রিন্টার থেকে স্লিপ বের হবে।",
        "স্লিপটি ছিঁড়ে সুন্দর করে রোগীর হাতে তুলে দিন।"
    ], "1-Click Printing")

    add_card(slide5, 8.8, 1.5, 3.7, 5.3, "প্রিন্টারে সমস্যা হলে সমাধান", [
        "সমস্যা ১: কাগজ বের হচ্ছে না? -> প্রিন্টারে পেপারের ঢাকনা ঠিকমতো লাগানো আছে কিনা দেখুন।",
        "সমস্যা ২: লেখা হালকা আসছে? -> থার্মাল পেপার উল্টো লাগানো থাকলে উল্টে দিন।",
        "সমস্যা ৩: লাল বাতি জ্বলছে? -> প্রিন্টারে কাগজ শেষ হয়ে গেছে, নতুন পেপার রোল লাগান।"
    ], "Troubleshooting", border_color=C_CRIMSON)

    # ==========================================
    # SLIDE 6: Step 5 - Cash & Digital Payment Collection
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "৫. বিল গ্রহণ: ক্যাশ, বিকাশ ও কার্ড (Billing & Payments)", "টাকা নেওয়া ও মানি রিসিট তৈরির নির্ভুল পদ্ধতি", "ধাপ ০৫")

    add_card(slide6, 0.8, 1.5, 3.6, 5.3, "ক্যাশ টাকা গ্রহণ (Cash)", [
        "রোগী ক্যাশ দিলে টাকা সুন্দর করে গুনে নিন।",
        "স্ক্রিনে 'Cash' অপশন সিলেক্ট করা থাকবে।",
        "প্রাপ্ত টাকার পরিমাণ লিখুন (যেমন ৳ ১,০০০ দিলে ফেরত ৳ ২০০)।",
        "সিস্টেম নিজেই দেখিয়ে দেবে ফেরত (Change) কত দিতে হবে!",
        "টাকা ক্যাশ ড্রয়ারে রেখে ড্রয়ার বন্ধ করুন।"
    ], "Cash Workflow")

    add_card(slide6, 4.8, 1.5, 3.6, 5.3, "বিকাশ ও ডিজিটাল পেমেন্ট", [
        "রোগী বিকাশ বা নগদ দিতে চাইলে 'bKash/Nagad' সিলেক্ট করুন।",
        "কাউন্টারের সামনে থাকা অফিসিয়াল QR কোড স্ক্যান করতে বলুন।",
        "পেমেন্ট সফল হলে রোগীর মোবাইল থেকে TrxID দেখে সফটওয়্যারে ইনপুট দিন।",
        "কার্ড পেমেন্ট হলে POS মেশিনের স্লিপের শেষ ৪ ডিজিট লিখুন।"
    ], "Digital Payment")

    add_card(slide6, 8.8, 1.5, 3.7, 5.3, "পেমেন্টের নির্ভুল নিয়মাবলী", [
        "নিয়ম ১: টাকা পাওয়ার আগে কখনো 'Paid' বাটনে চাপবেন না।",
        "নিয়ম ২: প্রতিটি পেমেন্টের সাথে সাথে অটো ইনভয়েস জেনারেট হবে।",
        "নিয়ম ৩: কোনো বাকি (Due) থাকলে তা রোগীর প্রোফাইলে স্বয়ংক্রিয় লাল কালিতে দেখাবে।",
        "মানি রিসিট রোগীকে বুঝিয়ে দিন।"
    ], "Zero-Error Rule", border_color=C_EMERALD)

    # ==========================================
    # SLIDE 7: Step 6 - Emergency Fast-Track
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "৬. জরুরি রোগী: ফাস্ট-ট্র্যাক ট্রায়াজ (Emergency Triage)", "সংকটাপন্ন রোগীর ক্ষেত্রে রিসেপশনের জীবন বাঁচানোর দায়িত্ব", "ধাপ ০৬")

    add_card(slide7, 0.8, 1.5, 3.6, 5.3, "কোন রোগীরা জরুরি (Emergency)?", [
        "১. সড়ক দুর্ঘটনা বা রক্তক্ষরণ হচ্ছে এমন রোগী।",
        "২. তীব্র বুকে ব্যথা / হার্ট অ্যাটাক বা স্ট্রোকের লক্ষণ।",
        "৩. তীব্র শ্বাসকষ্ট বা অজ্ঞান রোগী।",
        "৪. বিষপান বা আগুনে পোড়া রোগী।"
    ], "Emergency Criteria", border_color=C_CRIMSON)

    add_card(slide7, 4.8, 1.5, 3.6, 5.3, "জিরো-সেকেন্ড বিলম্ব নীতি (Zero Delay)", [
        "সবার আগে রোগীকে ইমার্জেন্সি রুমে ট্রলি বা হুইলচেয়ারে পাঠান!",
        "টাকা বা রেজিস্ট্রেশনের জন্য ১ সেকেন্ডও রোগীকে আটকে রাখবেন না।",
        "ডাক্তার ও নার্সকে সাথে সাথে খবর দিন।",
        "জীবন বাঁচানোই হাসপাতালের প্রথম ও প্রধান কাজ।"
    ], "Life First Rule")

    add_card(slide7, 8.8, 1.5, 3.7, 5.3, "ইমার্জেন্সি রেজিস্ট্রেশন কীভাবে করবেন", [
        "স্ক্রিনে থাকা লাল 'EMERGENCY FAST-TRACK' বাটনে চাপুন (বা F4 চাপুন)।",
        "রোগীর আত্মীয়কে বলুন পরে নাম ও ঠিকানা দিয়ে যেতে।",
        "সিস্টেম একটি তাৎক্ষণিক ইমার্জেন্সি আইডি খুলে ট্র্যাকিং চালু করবে।",
        "রোগী বিপদমুক্ত হওয়ার পর স্বাভাবিক বিলিং সম্পন্ন করুন।"
    ], "Fast-Track Protocol", border_color=C_ACCENT)

    # ==========================================
    # SLIDE 8: Step 7 - IPD Admission & Bed Allocation
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "৭. ভর্তি রোগী ও সিট/কেবিন বরাদ্দ (IPD Bed Allocation)", "রোগী হাসপাতালে ভর্তি হলে কেবিন বা সিট দেওয়ার সহজ নিয়ম", "ধাপ ০৭")

    add_card(slide8, 0.8, 1.5, 3.6, 5.3, "ধাপ ১: বেড সিলেকশন ম্যাট্রিক্স", [
        "'IPD Admission' মেনুতে ক্লিক করুন।",
        "স্ক্রিনে পুরো হাসপাতালের সব রুমের কালার-কোড ম্যাপ দেখতে পাবেন:",
        "সবুজ কালার = বেড সম্পূর্ণ খালি (Ready for patient)।",
        "লাল কালার = বেড ইতিমধ্যে ভর্তি (Occupied)।",
        "হলুদ কালার = রোগীকে রিলিজ দেওয়া হয়েছে, বেড পরিষ্কার হচ্ছে।"
    ], "Color Matrix")

    add_card(slide8, 4.8, 1.5, 3.6, 5.3, "ধাপ ২: রোগীর পছন্দ ও সামর্থ্য", [
        "রোগী বা তার পরিবারকে অপশনগুলো সুন্দর করে বুঝিয়ে বলুন:",
        "১. জেনারেল ওয়ার্ড (পুরুষ / মহিলা আলাদা)।",
        "২. সেমি-কেবিন / এসি নন-এসি কেবিন।",
        "৩. ভিআইপি ডিলাক্স কেবিন বা আইসিইউ (ICU/CCU)।",
        "প্রতিদিনের ভাড়া ও সুবিধাসমূহ বুঝিয়ে বলুন।"
    ], "Cabin Categories")

    add_card(slide8, 8.8, 1.5, 3.7, 5.3, "ধাপ ৩: অগ্রিম ডিপোজিট ও ভর্তি ফর্ম", [
        "খালি বেডের উপর ক্লিক করলেই রোগীর নামে বুকিং হয়ে যাবে।",
        "ভর্তি ফি / অ্যাডভান্স ডিপোজিট সংগ্রহ করুন (যেমন ৳ ৫,০০০ বা ৳ ১০,০০০)।",
        "ডিপোজিটের মানি রিসিট রোগীকে দিন।",
        "ওয়ার্ড বয়কে রোগীর ফাইলসহ নির্দিষ্ট কেবিনে যাওয়ার নির্দেশ দিন।"
    ], "Admission Deposit", border_color=C_SECONDARY)

    # ==========================================
    # SLIDE 9: Step 8 - Mistake Correction & Void
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "৮. ভুল হলে কি করবেন? (Mistake Correction & Void)", "ভয় পাওয়ার কিছু নেই! সব ভুলের নিরাপদ সমাধানের নিয়ম", "ধাপ ০৮")

    add_card(slide9, 0.8, 1.5, 3.6, 5.3, "কেস ১: নাম বা বয়স ভুল হয়েছে", [
        "কোনো চিন্তা নেই! রোগীর প্রোফাইল ওপেন করুন।",
        "ডানপাশের 'Edit Patient' পেনসিল আইকনে ক্লিক করুন।",
        "সঠিক বানান বা বয়স লিখে 'Update' চাপুন।",
        "মুহূর্তের মধ্যে সঠিক তথ্য সেভ হয়ে যাবে।"
    ], "Profile Edit")

    add_card(slide9, 4.8, 1.5, 3.6, 5.3, "কেস ২: ভুল ডাক্তার বা টেস্ট সিলেক্ট হয়েছে", [
        "যদি ভুল কোনো সার্ভিসের বিল হয়ে যায়:",
        "ইনভয়েসের পাশে থাকা 'Void Service' বাটন চাপুন।",
        "কারণ লিখুন (যেমন: 'রোগী অন্য ডাক্তারের কাছে যাবেন')।",
        "ভুল সার্ভিসটি বাতিল হয়ে সঠিক অংক হিসাবভুক্ত হবে।",
        "সিস্টেমের একাউন্টিং সবসময় নিখুঁত ব্যালেন্স রাখবে।"
    ], "Safe Void")

    add_card(slide9, 8.8, 1.5, 3.7, 5.3, "যা কখনো করবেন না (সতর্কতা)", [
        "সতর্কতা ১: রোগীর আসল ডেটা না জেনে ডিলিট করার চেষ্টা করবেন না।",
        "সতর্কতা ২: টাকা ফেরত দিলে সিস্টেমে অবশ্যই 'Refund' এন্ট্রি করবেন।",
        "সতর্কতা ৩: কোনো সমস্যা আটকে গেলে সরাসরি সুপারভাইজারকে জানান।",
        "সিস্টেমের সবকিছু অডিট লগে সুরক্ষিত থাকে।"
    ], "Strict Warning", border_color=C_CRIMSON)

    # ==========================================
    # SLIDE 10: Step 9 - Shift End & Cash Reconciliation
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "৯. শিফট সমাপ্তি ও ক্যাশ মেলানো (Shift Closing & Cash)", "ডিউটি শেষে নির্ভুলভাবে টাকা গুনে পরবর্তী শিফটে হ্যান্ডওভার করার নিয়ম", "ধাপ ০৯")

    add_card(slide10, 0.8, 1.5, 3.6, 5.3, "ধাপ ১: শিফট ক্লোজিং স্ক্রিন ওপেন", [
        "ডিউটি শেষে ড্যাশবোর্ডের উপরে 'End Shift' বাটনে ক্লিক করুন।",
        "স্ক্রিনে সারাদিনের মোট কালেকশন দেখতে পাবেন:",
        "- মোট ক্যাশ কালেকশন",
        "- মোট বিকাশ / নগদ কালেকশন",
        "- মোট কার্ড কালেকশন",
        "- মোট রিফান্ড ও নেট ক্যাশ ব্যালেন্স।"
    ], "Shift Screen")

    add_card(slide10, 4.8, 1.5, 3.6, 5.3, "ধাপ ২: ক্যাশ টাকা গুনে এন্ট্রি দিন", [
        "ক্যাশ ড্রয়ার খুলে সব টাকা সামনে রাখুন।",
        "১০০০, ৫০০, ২০০, ১০০ টাকার নোটগুলো গুনে সফটওয়্যারে সংখ্যা লিখুন।",
        "সফটওয়্যার স্বয়ংক্রিয়ভাবে দেখাবে: Difference = ৳ 0.00 (Perfect Match!)",
        "যদি পার্থক্য থাকে তাহলে কোনো ভাউচার বা খরচ বাদ পড়েছে কিনা চেক করুন।"
    ], "Cash Denomination")

    add_card(slide10, 8.8, 1.5, 3.7, 5.3, "ধাপ ৩: রিপোর্ট প্রিন্ট ও হ্যান্ডওভার", [
        "'Print Shift Summary' বাটনে ক্লিক করে শিফট রিপোর্ট প্রিন্ট করুন।",
        "রিপোর্টে নিজের স্বাক্ষর করুন।",
        "পরবর্তী রিসেপশনিস্টের হাতে টাকা ও প্রিন্ট করা কপি বুঝিয়ে দিয়ে স্বাক্ষর নিন।",
        "লগআউট করে নিশ্চিন্তে বাড়ি যান!"
    ], "Handover Complete", border_color=C_EMERALD)

    # ==========================================
    # SLIDE 11: Step 10 - Keyboard Shortcuts & Fast Operations
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "১০. দ্রুত কাজের জন্য কীবোর্ড শর্টকাট (Keyboard Shortcuts)", "মাউস ছাড়া বিদ্যুৎ গতিতে কাজ করার জাদুকরী কি-বোর্ড শর্টকাট", "শর্টকাট")

    add_card(slide11, 0.8, 1.5, 3.6, 5.3, "প্রধান শর্টকাট তালিকা", [
        "Alt + N : নতুন রোগী নিবন্ধন (New Patient)",
        "Alt + B : নতুন বিলিং স্ক্রিন (New Billing)",
        "Alt + D : ডাক্তার তালিকা ও শিডিউল (Doctors)",
        "Ctrl + P : সরাসরি থার্মাল প্রিন্ট (Instant Print)",
        "F2 : দ্রুত রোগী খোঁজা (Global Patient Search)",
        "F4 : ইমার্জেন্সি রেড অ্যালার্ট (Emergency Triage)",
        "Esc : যেকোনো উইন্ডো বা পপআপ বন্ধ করা"
    ], "Fast Keys", border_color=C_PRIMARY)

    add_card(slide11, 4.8, 1.5, 3.6, 5.3, "রোগী খোঁজার জাদুকরী টিপস", [
        "রোগী আগে কখনো আসলে তার মোবাইল নম্বরের শেষ ৪ ডিজিট লিখলেই চলে আসবে।",
        "অথবা রোগীর নাম লিখে এন্টার দিলেই পূর্বের সমস্ত হিস্ট্রি চলে আসবে।",
        "একই রোগীর জন্য কখনোই দ্বিতীয়বার নতুন আইডি বানাবেন না!",
        "পুরনো আইডিতেই নতুন ভিজিট যুক্ত করুন।"
    ], "Smart Search")

    add_card(slide11, 8.8, 1.5, 3.7, 5.3, "রিসেপশনিস্ট এটিকেট ও ব্যবহার", [
        "১. রোগী আসলে প্রথমে হাসিমুখে সালাম বা নমস্কার দিন।",
        "২. নরম ও আন্তরিক গলায় কথা বলুন; অসুস্থ রোগী ও স্বজনরা অস্থির থাকতে পারেন।",
        "৩. দিকনির্দেশনা স্পষ্ট করে দিন (ডাক্তার কোন তলায় ও কোন রুমে)।",
        "আপনার একটি মিষ্টি হাসিতে রোগীর অর্ধেক কষ্ট কমে যায়!"
    ], "Hospitality Best Practice", border_color=C_SECONDARY)

    # ==========================================
    # SLIDE 12: Step 11 - 24/7 IT Support & Emergency Contacts
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "১১. ২৪/৭ আইটি সহায়তা ও হেল্পলাইন (IT Support & Contacts)", "সিস্টেমে কোনো সমস্যা হলে সাথে সাথে যোগাযোগের নির্ভরযোগ্য মাধ্যম", "সহায়তা")

    add_card(slide12, 0.8, 1.5, 3.6, 5.3, "২৪/৭ টেলিগ্রাম অটোমেশন সাপোর্ট", [
        "যেকোনো টেকনিক্যাল প্রশ্ন বা রিপোর্টে টেলিগ্রাম ব্যবহার করুন।",
        "অফিসিয়াল টেলিগ্রাম বট: @onnesha_hospital_bot",
        "টেলিগ্রামে যেকোনো স্ট্যাটাস বা জরুরি অ্যালার্ট সরাসরি চেক করা যায়।",
        "আপনার মোবাইল থেকেও সম্পূর্ণ সিস্টেম মনিটর করা সম্ভব।"
    ], "Telegram Bot Cloud")

    add_card(slide12, 4.8, 1.5, 3.6, 5.3, "জরুরি হাসপাতাল হটলাইন নম্বর", [
        "হাসপাতাল এডমিন ও ম্যানেজার ফোন: ০১৭১X-XXXXXX",
        "আইটি সাপোর্ট ইনচার্জ: ০১৭XX-XXXXXX",
        "ইমার্জেন্সি ও অ্যাম্বুলেন্স হটলাইন: ০১৯XX-XXXXXX",
        "ফার্মেসি ও প্যাথলজি ল্যাব এক্সটেনশন: ১০২, ১০৫",
        "ডাক্তার ডিউটি রুম এক্সটেনশন: ২০১"
    ], "Emergency Hotlines")

    add_card(slide12, 8.8, 1.5, 3.7, 5.3, "সিস্টেমের অবস্থান ও নিরাপত্তা", [
        "সিস্টেম সংস্করণ: v1.1.90 (Latest Production)",
        "ক্লাউড প্ল্যাটফর্ম: Cloudflare Pages Edge Network",
        "ডাটাবেস: Supabase Enterprise Database (Auto-Backup)",
        "আপনার প্রতি সেকেন্ডের কাজ ক্লাউডে ১০০% সুরক্ষিত!",
        "ধন্যবাদ ও শুভ কর্মদিবস!"
    ], "Security & Cloud", border_color=C_EMERALD)

    prs.save(output_path)
    print(f"SUCCESS: Saved presentation to {output_path}")

if __name__ == "__main__":
    desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
    out_file = os.path.join(desktop_dir, "ONNESHA_HOSPITAL_RECEPTION_BABY_STEP_GUIDE.pptx")
    create_deck(out_file)
